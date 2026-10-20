import io
import logging
import re
import time
import zipfile
import zlib
from itertools import islice
from pathlib import Path

import docx
import pdfplumber
import uvicorn
from charset_normalizer import from_bytes
from fastapi import FastAPI, HTTPException, Request, UploadFile
from fastapi.responses import JSONResponse
from lxml.etree import XMLSyntaxError
from pdfminer.pdfpage import PDFPage

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

app = FastAPI()

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_REQUEST_SIZE = MAX_FILE_SIZE + 1024 * 1024
MAX_UNPACKED_SIZE = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_PDF_PAGES = 500
PARSE_TIME_LIMIT = 20
SUPPORTED_FORMATS = ("txt", "docx", "pdf")

# Дозволені резервні кодування для TXT
FALLBACK_ENCODINGS = {
    "cp1251", "koi8_r", "koi8_u", "cp866", "cp1125",
    "utf_16", "utf_16_le", "utf_16_be",
}  # fmt: skip

# Дієслова-накази для питань без знаку "?"
QUESTION_STARTERS = {
    "поясни", "поясніть", "розкажи", "розкажіть", "опиши", "опишіть",
    "назвіть", "наведи", "наведіть", "порівняй", "порівняйте",
    "визнач", "визначте", "перелічи", "перелічіть",
    "explain", "describe", "define", "compare", "discuss",
}  # fmt: skip

NUMBERING_RE = re.compile(r"^\s*\d+\s*[.)\-](?!\d)\s*")
FIRST_WORD_RE = re.compile(r"^\w+")
SENTENCE_END = (".", "?", "!", ":", ";")


def safe_for_log(value: str | None, limit: int = 100) -> str | None:
    """Очищує рядок від переносів та обрізає довжину для логування."""
    if value is None:
        return None
    return "".join(ch if ch.isprintable() else "?" for ch in value)[:limit]


def get_file_format(filename: str | None) -> str:
    """Розширення малими літерами без крапки. Порожній рядок, якщо його нема."""
    return Path(filename or "").suffix.lower().lstrip(".")


def decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")  # utf-8-sig прибирає BOM
    except UnicodeDecodeError:
        pass

    match = from_bytes(content, cp_isolation=list(FALLBACK_ENCODINGS | {"utf_8"})).best()
    encoding = match.encoding if match else None
    if encoding not in FALLBACK_ENCODINGS:
        raise HTTPException(
            status_code=400,
            detail="Не вдалося визначити кодування TXT файлу. Збережіть його у UTF-8",
        )
    try:
        return content.decode(encoding)
    except UnicodeDecodeError as exc:
        raise HTTPException(
            status_code=400,
            detail="Не вдалося прочитати TXT файл, можливо він пошкоджений",
        ) from exc


def parser_txt(content: bytes) -> str:
    return decode_text(content)


def check_pdf_pages(pdf) -> None:
    """Рахує сторінки з зупинкою на ліміті, а не обходить увесь документ."""
    pages = PDFPage.create_pages(pdf.doc)
    if sum(1 for _ in islice(pages, MAX_PDF_PAGES + 1)) > MAX_PDF_PAGES:
        raise HTTPException(status_code=413, detail="PDF має забагато сторінок")


def parser_pdf(content: bytes) -> str:
    deadline = time.monotonic() + PARSE_TIME_LIMIT
    pages = []
    try:
        pdf = pdfplumber.open(io.BytesIO(content))
        check_pdf_pages(pdf)
        with pdf:
            for page in pdf.pages:
                if time.monotonic() > deadline:
                    raise HTTPException(
                        status_code=413, detail="Обробка PDF зайняла забагато часу"
                    )
                page_text = page.extract_text()
                page.close()
                if page_text:
                    pages.append(page_text)
    except HTTPException:
        raise
    except Exception as exc:  # pdfminer на битих файлах кидає що завгодно
        raise HTTPException(
            status_code=400,
            detail="Не вдалося прочитати PDF файл, можливо він пошкоджений",
        ) from exc
    return "\n".join(pages)


def check_zip_size(content: bytes) -> None:
    """DOCX це zip. Дивимось на заявлені розміри ДО розпакування."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
    total_size = sum(entry.file_size for entry in entries)
    if len(entries) > MAX_ZIP_ENTRIES or total_size > MAX_UNPACKED_SIZE:
        raise HTTPException(
            status_code=413, detail="DOCX занадто великий після розпакування"
        )


def parser_docx(content: bytes) -> str:
    blocks = []
    try:
        check_zip_size(content)
        document = docx.Document(io.BytesIO(content))
        blocks.extend(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            seen = set()  # уникаємо дублювання об'єднаних комірок
            for row in table.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        blocks.append(cell.text)
    except (
        zipfile.BadZipFile,
        KeyError,
        XMLSyntaxError,
        zlib.error,
        ValueError,
        NotImplementedError,
    ) as exc:
        raise HTTPException(
            status_code=400,
            detail="Не вдалося прочитати DOCX файл, можливо він пошкоджений",
        ) from exc
    return "\n\n".join(blocks)


def tru_format(file_format: str, content: bytes) -> str:
    if file_format == "pdf":
        return parser_pdf(content)
    if file_format == "docx":
        return parser_docx(content)
    if file_format == "txt":
        return parser_txt(content)
    raise ValueError(f"Непідтримуваний формат файлу: {file_format!r}")


def merge_wrapped_lines(text: str) -> list[str]:
    """Склеює рядки, якщо наступний рядок є продовженням попереднього."""
    groups: list[list[str]] = []
    previous_blank = True
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            previous_blank = True
            continue
        is_continuation = (
            bool(groups)
            and not previous_blank
            and not groups[-1][-1].endswith(SENTENCE_END)
            and not NUMBERING_RE.match(line)
            and (
                line[0].islower()
                or (NUMBERING_RE.match(groups[-1][0]) and line.endswith("?"))
            )
        )
        if is_continuation:
            groups[-1].append(line)
        else:
            groups.append([line])
        previous_blank = False
    return [" ".join(parts) for parts in groups]


def is_question(line: str) -> bool:
    if line.endswith("?"):
        return True
    first_word = FIRST_WORD_RE.match(line)
    return bool(first_word) and first_word.group().lower() in QUESTION_STARTERS


def find_quetion(text: str) -> list[str]:
    questions = []
    for line in merge_wrapped_lines(text):
        cleaned = NUMBERING_RE.sub("", line).strip()
        if cleaned and is_question(cleaned):
            questions.append(cleaned)
    return questions


def handle_parse(file: UploadFile) -> dict:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Не вказано ім'я файлу")

    file_format = get_file_format(file.filename)
    if file_format not in SUPPORTED_FORMATS:
        raise HTTPException(status_code=400, detail="Невірний формат файлу")

    # Попередня перевірка розміру
    if file.size is not None and file.size > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="Файл завеликий")

    # Читання з контролем ліміту
    content = file.file.read(MAX_FILE_SIZE + 1)
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="Файл завеликий")
    if not content:
        raise HTTPException(status_code=400, detail="Файл порожній")

    questions = find_quetion(tru_format(file_format, content))
    return {
        "filename": file.filename,
        "source_format": file_format,
        "total": len(questions),
        "questions": questions,
    }


def log_request(
    filename: str | None,
    status_code: int,
    duration: float,
    error: Exception | None = None,
) -> None:
    if status_code >= 500:
        level = logging.ERROR
    elif status_code >= 400:
        level = logging.WARNING
    else:
        level = logging.INFO
    logger.log(
        level,
        "code=%s file=%s duration=%.2f",
        status_code,
        safe_for_log(filename),
        duration,
        exc_info=error,
    )


@app.middleware("http")
async def limit_request_size(request: Request, call_next):
    """Відсікає завеликий запит за Content-Length до парсингу тіла."""
    if request.method != "POST":
        return await call_next(request)

    start = time.perf_counter()
    length = request.headers.get("content-length")
    if length is None:
        status_code, detail = 411, "Потрібен заголовок Content-Length"
    elif not length.isdigit():
        status_code, detail = 400, "Некоректний заголовок Content-Length"
    elif int(length) > MAX_REQUEST_SIZE:
        status_code, detail = 413, "Файл завеликий"
    else:
        return await call_next(request)

    log_request(None, status_code, time.perf_counter() - start)
    return JSONResponse(status_code=status_code, content={"detail": detail})


# Синхронний ендпоінт автоматично виконується у threadpool
@app.post("/parse")
def parser_file(file: UploadFile):
    start = time.perf_counter()
    status_code = 200
    error: Exception | None = None
    try:
        return handle_parse(file)
    except HTTPException as exc:
        status_code = exc.status_code
        raise
    except Exception as exc:
        status_code = 500
        error = exc
        raise HTTPException(
            status_code=500, detail="Внутрішня помилка сервера"
        ) from exc
    finally:
        log_request(file.filename, status_code, time.perf_counter() - start, error)


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8001)