import io
import logging
import zipfile
from types import SimpleNamespace

import docx
import main
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from main import app, find_quetion, get_file_format, tru_format

client = TestClient(app)


# ---------- хелпери ----------


def post_file(filename: str, content: bytes):
    return client.post("/parse", files={"file": (filename, content)})


def questions_of(filename: str, content: bytes) -> list[str]:
    response = post_file(filename, content)
    assert response.status_code == 200, response.text
    return response.json()["questions"]


def make_docx(paragraphs=(), table_rows=()) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docm() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            b'<Override PartName="/word/document.xml" '
            b'ContentType="application/vnd.ms-word.document.macroEnabled.main+xml"/>'
            b"</Types>",
        )
        z.writestr(
            "_rels/.rels",
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Id="rId1" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            b'Target="word/document.xml"/>'
            b"</Relationships>",
        )
        z.writestr(
            "word/document.xml",
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>',
        )
    return buffer.getvalue()


def make_pdf(lines: list[str], pages: int = 1, rotate: str | None = None) -> bytes:
    """Мінімальний валідний PDF (тільки ASCII). Усі сторінки ділять один вміст."""
    commands = ["BT", "/F1 12 Tf", "14 TL", "50 750 Td"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"({escaped}) Tj T*")
    commands.append("ET")
    stream = "\n".join(commands).encode("latin-1")

    kids = " ".join(f"{5 + i} 0 R" for i in range(pages)).encode()
    rotate_str = f" /Rotate {rotate}" if rotate else ""
    page_object = (
        f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]{rotate_str} "
        "/Contents 4 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
    ).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % pages,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        *[page_object] * pages,
    ]
    pdf = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n" % number + body + b"\nendobj\n"

    xref_position = len(pdf)
    pdf += b"xref\n0 %d\n" % (len(objects) + 1)
    pdf += b"0000000000 65535 f \n"
    for offset in offsets:
        pdf += b"%010d 00000 n \n" % offset
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
    pdf += b"startxref\n%d\n%%%%EOF\n" % xref_position
    return pdf


# ---------- розпізнавання питань (TXT) ----------


def test_normal_questions_list():
    content = (
        "1. Що таке Python?\n"
        "2. Чим list відрізняється від tuple?\n"
        "3) Що таке декоратор?\n"
        "4 - Навіщо потрібен GIL?\n"
    ).encode()
    assert questions_of("questions.txt", content) == [
        "Що таке Python?",
        "Чим list відрізняється від tuple?",
        "Що таке декоратор?",
        "Навіщо потрібен GIL?",
    ]


def test_technical_names_with_dots_are_not_cut():
    content = (
        "1. What is Node.js?\n"
        "2. Поясни різницю між Promise та async/await в Node.js?\n"
        "3. Що робить asyncio.sleep()?\n"
    ).encode()
    assert questions_of("tech.txt", content) == [
        "What is Node.js?",
        "Поясни різницю між Promise та async/await в Node.js?",
        "Що робить asyncio.sleep()?",
    ]


def test_intro_text_without_question_mark_is_excluded():
    content = (
        "Це вступний текст без питань.\n"
        "Python - це мова програмування.\n"
        "1. А це вже справжнє питання?\n"
    ).encode()
    assert questions_of("mixed.txt", content) == ["А це вже справжнє питання?"]


def test_imperative_prompt_without_question_mark_is_question():
    # окремо від вступного тексту: файл містить тільки такі пункти
    content = (
        "1. Поясни dependency injection\n"
        "Розкажи про типи даних в Python.\n"
        "3. Поясни, як працює GIL\n"
        "4. Explain the event loop\n"
    ).encode()
    assert questions_of("prompts.txt", content) == [
        "Поясни dependency injection",
        "Розкажи про типи даних в Python.",
        "Поясни, як працює GIL",
        "Explain the event loop",
    ]


def test_numbered_statement_is_not_question():
    content = (
        "1. Python - це мова програмування\n2. Django - це веб-фреймворк.\n"
    ).encode()
    assert questions_of("statements.txt", content) == []


@pytest.mark.parametrize(
    "text",
    [
        "Do not use globals.",
        "List of topics:",
        "Name of the author: John",
        "Як зазначено вище, Python простий.",
        "Чи варто це робити. Скоріше ні.",
    ],
)
def test_ordinary_sentences_are_not_false_positives(text):
    assert find_quetion(text) == []


def test_multiline_question_is_merged():
    content = (
        "1. Поясни різницю між Promise\n"
        "та async/await в Node.js?\n"
        "2. Що таке декоратор?\n"
    ).encode()
    assert questions_of("wrapped.txt", content) == [
        "Поясни різницю між Promise та async/await в Node.js?",
        "Що таке декоратор?",
    ]


def test_headings_are_not_merged_with_neighbours():
    content = (
        "Звичайні питання:\n"
        "1. Що таке Python?\n"
        "Категорія\n"
        "Що таке декоратор?\n"
    ).encode()
    assert questions_of("headings.txt", content) == [
        "Що таке Python?",
        "Що таке декоратор?",
    ]


def test_uppercase_extension_is_accepted():
    assert questions_of("QUESTIONS.TXT", b"1. What is Python?\n") == [
        "What is Python?"
    ]


def test_response_contains_total_and_source_format():
    response = post_file("q.txt", b"1. What is A?\n2. What is B?\n")
    assert response.status_code == 200
    assert response.json() == {
        "filename": "q.txt",
        "source_format": "txt",
        "total": 2,
        "questions": ["What is A?", "What is B?"],
    }


# ---------- кодування TXT ----------

CYRILLIC_TEXT = (
    "1. Що таке Python?\n"
    "2. Чим list відрізняється від tuple?\n"
    "3. Що таке декоратор і навіщо він потрібен?\n"
    "4. Чим відрізняється процес від потоку?\n"
    "5. Що таке GIL і як він впливає на багатопотоковість?\n"
    "6. Поясни dependency injection\n"
    "7. Що таке генератор в Python?\n"
    "8. Як працює збирач сміття?\n"
)


def test_txt_cp1251_is_decoded():
    result = questions_of("cp1251.txt", CYRILLIC_TEXT.encode("cp1251"))
    assert result[0] == "Що таке Python?"
    assert len(result) == 8


def test_txt_utf8_with_bom_is_decoded():
    content = b"\xef\xbb\xbf" + CYRILLIC_TEXT.encode()
    result = questions_of("bom.txt", content)
    assert result[0] == "Що таке Python?"


def test_txt_utf16_is_decoded():
    result = questions_of("utf16.txt", CYRILLIC_TEXT.encode("utf-16"))
    assert len(result) == 8


def test_txt_with_unknown_encoding_returns_400_not_garbage():
    junk = bytes(range(128, 256)) * 3
    response = post_file("junk.txt", junk)
    assert response.status_code == 400
    assert "кодування" in response.json()["detail"]


# ---------- PDF і DOCX ----------


def test_valid_pdf_is_parsed():
    pdf = make_pdf(
        ["Intro text without question.", "1. What is Python?", "2. Explain GIL"]
    )
    response = post_file("questions.pdf", pdf)
    assert response.status_code == 200
    body = response.json()
    assert body["source_format"] == "pdf"
    assert body["questions"] == ["What is Python?", "Explain GIL"]


def test_valid_docx_is_parsed():
    content = make_docx(["Вступ:", "1. Що таке Python?", "2. Поясни декоратори"])
    assert questions_of("questions.docx", content) == [
        "Що таке Python?",
        "Поясни декоратори",
    ]


def test_docx_table_questions_are_read():
    content = make_docx(
        paragraphs=["Список питань"],
        table_rows=[
            ["Питання", "Категорія"],
            ["Що таке декоратор?", "Python"],
            ["Поясни dependency injection", "Архітектура"],
        ],
    )
    assert questions_of("table.docx", content) == [
        "Що таке декоратор?",
        "Поясни dependency injection",
    ]


# ---------- помилки валідації ----------


@pytest.mark.parametrize(
    "filename", ["virus.exe", "questions", "archive.tar.gz"]
)
def test_invalid_format_returns_400(filename):
    response = post_file(filename, b"some content")
    assert response.status_code == 400
    assert response.json()["detail"] == "Невірний формат файлу"


def test_empty_file_returns_400():
    assert post_file("empty.txt", b"").status_code == 400


def test_max_file_size_is_10_mb():
    assert main.MAX_FILE_SIZE == 10 * 1024 * 1024


# Для швидкості (multipart-парсер повільний) перевіряємо логіку ліміту на малому
# значенні: сам код не залежить від того, чи ліміт 1 КБ чи 10 МБ.
def test_file_too_large_returns_413(monkeypatch):
    monkeypatch.setattr(main, "MAX_FILE_SIZE", 1024)
    assert post_file("big.txt", b"a" * 1025).status_code == 413


def test_file_of_exact_limit_is_accepted(monkeypatch):
    monkeypatch.setattr(main, "MAX_FILE_SIZE", 1024)
    response = post_file("limit.txt", b"a" * 1024)
    assert response.status_code == 200
    assert response.json()["total"] == 0


def test_too_large_file_is_rejected_by_size_before_reading():
    class ExplodingFile:
        def read(self, *args):
            raise AssertionError("файл не повинен читатись")

    upload = SimpleNamespace(
        filename="big.txt", size=main.MAX_FILE_SIZE + 1, file=ExplodingFile()
    )
    with pytest.raises(HTTPException) as error:
        main.handle_parse(upload)
    assert error.value.status_code == 413


def test_filename_none_returns_400_instead_of_crash():
    upload = SimpleNamespace(filename=None, size=1, file=io.BytesIO(b"x"))
    with pytest.raises(HTTPException) as error:
        main.handle_parse(upload)
    assert error.value.status_code == 400
    assert "ім'я файлу" in error.value.detail


# ---------- пошкоджені файли ----------


def broken_docx_xml() -> bytes:
    source = zipfile.ZipFile(io.BytesIO(make_docx(["text"])))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name in source.namelist():
            data = b"<broken" if name == "word/document.xml" else source.read(name)
            archive.writestr(name, data)
    return output.getvalue()


def plain_zip() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("random.txt", "hello")
    return output.getvalue()


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        pytest.param("broken.pdf", b"this is not a real pdf", id="not-a-pdf"),
        pytest.param("broken.docx", b"this is not a real docx", id="not-a-zip"),
        pytest.param("renamed.docx", plain_zip(), id="zip-without-docx-parts"),
        pytest.param("broken_xml.docx", broken_docx_xml(), id="broken-xml"),
        pytest.param("truncated.docx", make_docx(["text"])[:300], id="truncated"),
    ],
)
def test_corrupted_files_return_400(filename, content):
    response = post_file(filename, content)
    assert response.status_code == 400
    assert "пошкоджений" in response.json()["detail"]


# ---------- tru_format ----------


def test_unknown_format_raises_value_error():
    with pytest.raises(ValueError):
        tru_format("exe", b"data")


# ---------- логування ----------


def test_client_error_is_logged_as_warning(caplog):
    with caplog.at_level(logging.INFO):
        post_file("virus.exe", b"data")
    records = [r for r in caplog.records if "code=400" in r.getMessage()]
    assert len(records) == 1
    assert records[0].levelno == logging.WARNING
    assert "virus.exe" in records[0].getMessage()


def test_success_is_logged_once_as_info(caplog):
    with caplog.at_level(logging.INFO):
        post_file("ok.txt", b"1. What is Python?\n")
    records = [r for r in caplog.records if "code=200" in r.getMessage()]
    assert len(records) == 1
    assert records[0].levelno == logging.INFO


def test_unexpected_error_returns_500_and_logs_traceback(monkeypatch, caplog):
    def boom(*args):
        raise RuntimeError("boom")

    monkeypatch.setattr(main, "tru_format", boom)
    with caplog.at_level(logging.INFO):
        response = post_file("ok.txt", b"1. What is Python?\n")
    assert response.status_code == 500
    assert "Traceback" in caplog.text
    assert "RuntimeError: boom" in caplog.text


# ---------- допоміжні функції ----------


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        (None, ""),
        ("", ""),
        ("questions", ""),
        ("Q.PDF", "pdf"),
        ("my.questions.Txt", "txt"),
    ],
)
def test_get_file_format(filename, expected):
    assert get_file_format(filename) == expected


# ---------- регресії (баги, знайдені під час рев'ю) ----------


def docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_table_cells_are_not_merged_into_one_question():
    content = make_docx(
        table_rows=[
            ["Поясни dependency injection", "архітектура"],
            ["Поясни GIL", "python"],
        ]
    )
    assert questions_of("table.docx", content) == [
        "Поясни dependency injection",
        "Поясни GIL",
    ]


def test_horizontally_merged_cell_gives_one_question():
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Що таке Python?"
    assert questions_of("merged.docx", docx_bytes(document)) == ["Що таке Python?"]


def test_vertically_merged_cell_gives_one_question():
    document = docx.Document()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Що таке Python?"
    assert questions_of("merged.docx", docx_bytes(document)) == ["Що таке Python?"]


def test_name_as_noun_is_not_question():
    assert find_quetion("Назви змінних мають бути зрозумілими.") == []


# ---------- безпека завантаження (issue #33) ----------


def test_oversized_request_is_rejected_before_parsing(monkeypatch):
    monkeypatch.setattr(main, "MAX_REQUEST_SIZE", 500)

    def must_not_run(*args):
        raise AssertionError("запит відхиляється ще до розбору тіла")

    monkeypatch.setattr(main, "handle_parse", must_not_run)
    response = post_file("big.txt", b"a" * 2000)
    assert response.status_code == 413


def test_request_without_content_length_is_rejected():
    def chunks():  # генератор змушує httpx слати chunked, без Content-Length
        yield b"--x\r\n"
        yield b"--x--\r\n"

    response = client.post(
        "/parse",
        content=chunks(),
        headers={"content-type": "multipart/form-data; boundary=x"},
    )
    assert response.status_code == 411


def make_zip_bomb(unpacked_size: int) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * unpacked_size)
    return buffer.getvalue()


def test_docx_zip_bomb_is_rejected():
    bomb = make_zip_bomb(main.MAX_UNPACKED_SIZE + 1)
    assert len(bomb) < 200 * 1024  # "бомба" справді мала: десятки КБ -> 50+ МБ
    response = post_file("bomb.docx", bomb)
    assert response.status_code == 413
    assert "після розпакування" in response.json()["detail"]


def test_docx_with_too_many_entries_is_rejected(monkeypatch):
    monkeypatch.setattr(main, "MAX_ZIP_ENTRIES", 3)
    assert post_file("many.docx", make_docx(["text"])).status_code == 413


def test_pdf_with_too_many_pages_is_rejected():
    pdf = make_pdf(["1. What is Python?"], pages=main.MAX_PDF_PAGES + 1)
    response = post_file("heavy.pdf", pdf)
    assert response.status_code == 413
    assert "сторінок" in response.json()["detail"]


def test_pdf_at_page_limit_is_accepted(monkeypatch):
    monkeypatch.setattr(main, "MAX_PDF_PAGES", 3)
    pdf = make_pdf(["1. What is Python?"], pages=3)
    assert post_file("ok.pdf", pdf).status_code == 200


def test_pdf_parse_time_limit_is_enforced(monkeypatch):
    monkeypatch.setattr(main, "PARSE_TIME_LIMIT", -1)  # дедлайн уже минув
    response = post_file("slow.pdf", make_pdf(["1. What is Python?"], pages=2))
    assert response.status_code == 413
    assert "забагато часу" in response.json()["detail"]


def test_newlines_in_filename_cannot_forge_log_lines(caplog):
    evil = "a.txt\nINFO | code=200 file=fake duration=0.00"
    with caplog.at_level(logging.INFO):
        main.log_request(evil, 400, 0.0)
    assert len(caplog.records) == 1
    assert "\n" not in caplog.records[0].getMessage()


def test_long_filename_is_truncated_in_logs():
    assert len(main.safe_for_log("a" * 500)) == 100
    assert main.safe_for_log(None) is None


def test_large_txt_linear_complexity():
    """200 000 рядків обробляються швидко без квадратичної складності."""
    import time

    text = "\n".join(["abc def"] * 200_000)
    start = time.perf_counter()
    lines = main.merge_wrapped_lines(text)
    duration = time.perf_counter() - start
    assert len(lines) == 1
    assert duration < 2.0


def test_docm_renamed_to_docx_returns_400():
    response = post_file("macro.docx", make_docm())
    assert response.status_code == 400
    assert "пошкоджений" in response.json()["detail"]


def test_pdf_with_rotate_abc_returns_400():
    pdf = make_pdf(["1. What is Python?"], rotate="/abc")
    response = post_file("corrupted_rotate.pdf", pdf)
    assert response.status_code == 400
    assert "пошкоджений" in response.json()["detail"]


def test_multiline_question_with_uppercase_continuation_is_merged():
    text = "1. What is the difference between\nPython 2 and Python 3?"
    questions = find_quetion(text)
    assert questions == ["What is the difference between Python 2 and Python 3?"]


def test_numbering_regex_does_not_strip_floats():
    text = "1.5 GB - це ліміт?"
    questions = find_quetion(text)
    assert questions == ["1.5 GB - це ліміт?"]
