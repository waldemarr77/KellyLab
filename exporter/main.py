import html
import io
import logging
import os
import re
import time
from urllib.parse import quote

import docx
import uvicorn
from fastapi import FastAPI, HTTPException, Request, Response
from fastapi.responses import JSONResponse
from pydantic import BaseModel, Field, field_validator
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

app = FastAPI(title="KellyLab Exporter Service")

EXPORT_TIME_LIMIT = 60
MAX_REQUEST_SIZE = 5 * 1024 * 1024  # 5 MB: 50 × (2000 + 20000) символів кирилицею ≈ 2.1 MB

CANDIDATE_FONTS = [
    os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf"),
    os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "arial.ttf"),
    os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "calibri.ttf"),
    os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "segoeui.ttf"),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/opentype/dejavu/DejaVuSans.ttf",
]

FONT_NAME = "AppCyrillicFont"
font_registered = False

for font_path in CANDIDATE_FONTS:
    if os.path.exists(font_path):
        try:
            pdfmetrics.registerFont(TTFont(FONT_NAME, font_path))
            font_registered = True
            logger.info("Успішно зареєстровано шрифт для PDF: %s", font_path)
            break
        except Exception as exc:  # noqa: BLE001 -- битий шрифт не має валити старт, пробуємо наступний
            logger.warning("Не вдалося завантажити шрифт %s: %s", font_path, exc)

if not font_registered:
    logger.error("Не знайдено жодного придатного шрифту з підтримкою кирилиці. PDF-експорт буде недоступний.")


# Керуючі символи, несумісні з XML (Word падає при їх записі)
_CONTROL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def strip_control_chars(text: str) -> str:
    """Видаляє керуючі символи, що ламають python-docx (ValueError: not XML compatible)."""
    return _CONTROL_CHARS_RE.sub("", text)


def check_font_available() -> None:
    """Fail-fast: якщо шрифт не знайдено, не створюємо пошкоджений файл."""
    if not font_registered:
        raise HTTPException(
            status_code=500,
            detail="Внутрішня помилка сервера: шрифт із кирилицею недоступний",
        )


class QuestionItem(BaseModel):
    order: int
    text: str = Field(..., min_length=1, max_length=2000, description="Текст питання (до 2000 символів)")
    answer: str = Field(default="", max_length=20000, description="Відповідь на питання (до 20000 символів)")

    @field_validator("answer", mode="before")
    @classmethod
    def coerce_answer(cls, v: object) -> str:
        if v is None:
            return ""
        return str(v)


class ExportRequest(BaseModel):
    format: str = Field(..., description="Формат файлу: 'pdf' або 'docx'")
    title: str = Field(..., min_length=1, max_length=255, description="Назва документа")
    questions: list[QuestionItem] = Field(
        default_factory=list,
        max_length=50,
        description="Список питань (до 50 елементів)",
    )

    @field_validator("format")
    @classmethod
    def validate_format(cls, v: str) -> str:
        cleaned = v.strip().lower()
        if cleaned not in ("pdf", "docx"):
            raise ValueError(f"Непідтримуваний формат: {v!r}. Дозволені: 'pdf', 'docx'")
        return cleaned

    @field_validator("title")
    @classmethod
    def validate_title(cls, v: str) -> str:
        cleaned = v.strip()
        if not cleaned:
            raise ValueError("Назва документа не може бути порожньою")
        return cleaned


def generate_docx(title: str, questions: list[QuestionItem]) -> bytes:
    """Генерація DOCX документа з підтримкою кирилиці."""
    doc = docx.Document()
    doc.add_heading(strip_control_chars(title), level=1)
    doc.add_paragraph("")

    for item in questions:
        p_q = doc.add_paragraph()
        run_q = p_q.add_run(f"{item.order}. {strip_control_chars(item.text)}")
        run_q.bold = True

        if item.answer.strip():
            p_a = doc.add_paragraph()
            p_a.paragraph_format.left_indent = docx.shared.Pt(18)
            p_a.add_run("Відповідь: ").italic = True
            p_a.add_run(strip_control_chars(item.answer))

        doc.add_paragraph("")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_pdf(title: str, questions: list[QuestionItem]) -> bytes:
    """Генерація PDF документа з підтримкою кирилиці через ReportLab."""
    check_font_available()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PDFTitle",
        parent=styles["Heading1"],
        fontName=FONT_NAME,
        fontSize=18,
        leading=22,
        spaceAfter=14,
        wordWrap="CJK",
    )
    question_style = ParagraphStyle(
        "PDFQuestion",
        parent=styles["Normal"],
        fontName=FONT_NAME,
        fontSize=12,
        leading=16,
        spaceAfter=6,
        wordWrap="CJK",
    )
    answer_style = ParagraphStyle(
        "PDFAnswer",
        parent=styles["Normal"],
        fontName=FONT_NAME,
        fontSize=10,
        leading=14,
        leftIndent=15,
        spaceAfter=12,
        textColor=HexColor("#333333"),
        wordWrap="CJK",
    )

    clean_title = strip_control_chars(title)
    story = [
        Paragraph(html.escape(clean_title).replace("\n", "<br/>"), title_style),
        Spacer(1, 10),
    ]

    for item in questions:
        clean_text = strip_control_chars(item.text)
        q_text = f"<b>{item.order}.</b> " + html.escape(clean_text).replace("\n", "<br/>")
        story.append(Paragraph(q_text, question_style))

        clean_answer = strip_control_chars(item.answer)
        if clean_answer.strip():
            a_text = "<i>Відповідь:</i> " + html.escape(clean_answer).replace("\n", "<br/>")
            story.append(Paragraph(a_text, answer_style))
        else:
            story.append(Spacer(1, 8))

    doc.build(story)
    return buffer.getvalue()


@app.middleware("http")
async def limit_request_size(request: Request, call_next):
    """Обмежує розмір вхідного запиту для запобігання переповнення пам'яті."""
    if request.method == "POST":
        length = request.headers.get("content-length")
        if length is not None:
            try:
                if int(length) > MAX_REQUEST_SIZE:
                    return JSONResponse(
                        status_code=413,
                        content={"detail": "Розмір запиту перевищує ліміт"},
                    )
            except ValueError:
                return JSONResponse(
                    status_code=400,
                    content={"detail": "Некоректний заголовок Content-Length"},
                )
        else:
            # Запит без Content-Length (наприклад, chunked transfer): читаємо потоком з лімітом
            body = bytearray()
            async for chunk in request.stream():
                body.extend(chunk)
                if len(body) > MAX_REQUEST_SIZE:
                    return JSONResponse(
                        status_code=413,
                        content={"detail": "Розмір запиту перевищує ліміт"},
                    )
            request._body = bytes(body)
    return await call_next(request)


@app.get("/health")
def health_check():
    """Ендпоінт для перевірки готовності сервісу (healthcheck для Docker / Compose)."""
    return {"status": "ok"}


@app.post("/export")
def export_file(request: ExportRequest):
    start = time.perf_counter()
    file_format = request.format
    sorted_questions = sorted(request.questions, key=lambda q: q.order)

    try:
        if file_format == "pdf":
            content = generate_pdf(request.title, sorted_questions)
            media_type = "application/pdf"
            extension = "pdf"
        else:
            content = generate_docx(request.title, sorted_questions)
            media_type = (
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            extension = "docx"

        duration = time.perf_counter() - start
        if duration > EXPORT_TIME_LIMIT:
            logger.warning("Експорт тривав %.2f с (перевищив ліміт %d с)", duration, EXPORT_TIME_LIMIT)
            raise HTTPException(status_code=504, detail="Експорт перевищив ліміт часу")

        clean_title = (
            re.sub(r'[\\/*?:"<>|\x00-\x1f]', "_", strip_control_chars(request.title.strip()))[:60].strip()
            or "export"
        )
        safe_filename = quote(f"{clean_title}.{extension}")
        headers = {
            "Content-Disposition": f'attachment; filename="export.{extension}"; filename*=UTF-8\'\'{safe_filename}',
        }

        logger.info(
            "code=200 format=%s title=%s duration=%.2f",
            file_format,
            clean_title,
            duration,
        )
        return Response(content=content, media_type=media_type, headers=headers)

    except HTTPException:
        raise
    except Exception:
        logger.exception("code=500 format=%s", file_format)
        raise HTTPException(status_code=500, detail="Внутрішня помилка сервера")


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8002)
